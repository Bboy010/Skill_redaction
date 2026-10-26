#!/usr/bin/env python3
"""
audit_zotero.py : audit des citations Zotero d'un manuscrit Word (style numéroté [1], [2]…).

LECTURE SEULE : le .docx n'est jamais modifié. Fermez-le dans Word avant de lancer le script.

Contrôles :
  1. numéros tapés à la main (hors champs Zotero) ;
  2. texte affiché du champ différent de la valeur enregistrée par Zotero ;
  3. chaque article cité (titre / DOI dans les données du champ) → numéro de la bibliographie ;
     signale les appels qui affichent un AUTRE numéro (citations restées à l'ancien numéro) ;
  4. références orphelines (jamais citées) et manquantes (citées mais absentes) ;
  5. ordre de première citation croissant ;
  6. balises de travail restantes « [ZOTERO: … ] ».

Usage :
    python audit_zotero.py manuscrit.docx
    python audit_zotero.py manuscrit.docx --style-biblio Bibliography --json rapport.json

Prérequis : pip install python-docx
"""
import argparse
import json
import re
import unicodedata
import docx
from docx.oxml.ns import qn

W = lambda t: qn("w:" + t)


def normaliser(s):
    """minuscules, sans accents ni ponctuation : pour comparer des titres."""
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def numeros(texte):
    """'[3,5–7]' -> [3, 5, 6, 7]"""
    out = []
    texte = texte.replace("\\uc0\\u8211{}", "–")
    for tok in texte.strip().strip("[]").split(","):
        m = re.fullmatch(r"\s*(\d+)\s*[–-]\s*(\d+)\s*", tok)
        if m:
            out += list(range(int(m.group(1)), int(m.group(2)) + 1))
        elif tok.strip().isdigit():
            out.append(int(tok))
    return out


def champs(d):
    """Reconstruit chaque champ complexe : (indice paragraphe, instruction, texte affiché, contexte)."""
    res = []
    for i, p in enumerate(d.paragraphs):
        etat, instr, vu = None, [], []
        for r in p._p.iter(W("r")):
            fc = r.find(W("fldChar"))
            if fc is not None:
                t = fc.get(W("fldCharType"))
                if t == "begin":
                    etat, instr, vu = "instr", [], []
                elif t == "separate":
                    etat = "resultat"
                elif t == "end":
                    if instr:
                        res.append((i, "".join(instr), "".join(vu), p.text[:60]))
                    etat = None
                continue
            if etat == "instr":
                it = r.find(W("instrText"))
                if it is not None:
                    instr.append(it.text or "")
            elif etat == "resultat":
                vu.append("".join(t.text or "" for t in r.findall(W("t"))))
    return res


def tapes_main(d, style_biblio):
    """Numéros [n] écrits hors de tout champ."""
    trouves = []
    for i, p in enumerate(d.paragraphs):
        if p.style.name == style_biblio:
            continue
        prof, buf = 0, []
        for r in p._p.iter(W("r")):
            fc = r.find(W("fldChar"))
            if fc is not None:
                t = fc.get(W("fldCharType"))
                prof += 1 if t == "begin" else (-1 if t == "end" else 0)
                continue
            if prof == 0:
                buf.append("".join(t.text or "" for t in r.findall(W("t"))))
        for m in re.finditer(r"\[\d[\d,–\- ]*\]", "".join(buf)):
            trouves.append((i, m.group(0), p.text[:60]))
    return trouves


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("--style-biblio", default="Bibliography", help="nom du style Word des entrées de bibliographie")
    ap.add_argument("--json", help="écrire le rapport dans ce fichier JSON")
    a = ap.parse_args()

    d = docx.Document(a.docx)
    tous = champs(d)
    cits = [c for c in tous if "ZOTERO_ITEM" in c[1]]
    biblio = {int(m.group(1)): p.text for p in d.paragraphs if p.style.name == a.style_biblio
              for m in [re.match(r"\s*\[?(\d+)[\].]\s", p.text)] if m}
    problemes, vus, ordre = [], set(), []

    for i, instr, vu, ctx in cits:
        js = json.loads(instr[instr.index("{"): instr.rindex("}") + 1])
        enregistre = js.get("properties", {}).get("plainCitation", "")
        if vu.strip() != enregistre.strip():
            problemes.append({"type": "affiché ≠ enregistré", "paragraphe": i, "affiché": vu, "enregistré": enregistre, "contexte": ctx})
        ns = numeros(enregistre)
        for item in js.get("citationItems", []):
            data = item.get("itemData", {})
            titre, doi = normaliser(data.get("title", ""))[:40], (data.get("DOI") or "").lower()
            cand = [k for k, v in biblio.items() if (doi and doi in v.lower()) or (titre and titre in normaliser(v))]
            if not cand:
                problemes.append({"type": "article absent de la bibliographie", "paragraphe": i, "titre": data.get("title", "")[:60], "contexte": ctx})
            elif not set(cand) & set(ns):
                problemes.append({"type": "appel → autre numéro (ancien numéro ?)", "paragraphe": i,
                                  "affiché": enregistre, "entrée attendue": cand, "titre": data.get("title", "")[:60], "contexte": ctx})
        for x in ns:
            if x not in vus:
                vus.add(x); ordre.append(x)

    tapes = tapes_main(d, a.style_biblio)
    balises = sum(p.text.count("[ZOTERO:") for p in d.paragraphs)
    orphelines = sorted(set(biblio) - vus)
    manquantes = sorted(vus - set(biblio))
    continue_ = sorted(biblio) == list(range(1, len(biblio) + 1))

    print(f"Champs de citation Zotero : {len(cits)} | bibliographie : {len(biblio)} entrées (numérotation continue : {continue_})")
    print(f"Numéros distincts cités : {len(vus)} | orphelines : {orphelines} | manquantes : {manquantes}")
    print(f"Ordre de première citation croissant : {ordre == sorted(ordre)}")
    print(f"Numéros tapés à la main : {len(tapes)} | balises [ZOTERO: …] restantes : {balises}")
    for t in tapes:
        print(f"  tapé à la main | paragraphe {t[0]} | {t[1]} | « {t[2]} »")
    print(f"Problèmes : {len(problemes)}")
    for pb in problemes:
        print("  " + " | ".join(f"{k}: {v}" for k, v in pb.items()))
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump({"champs": len(cits), "biblio": len(biblio), "orphelines": orphelines, "manquantes": manquantes,
                       "ordre_croissant": ordre == sorted(ordre), "tapes_main": tapes, "balises": balises,
                       "problemes": problemes}, f, ensure_ascii=False, indent=2)
        print(f"Rapport JSON : {a.json}")
