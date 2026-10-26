#!/usr/bin/env python3
"""
comparer_docx.py : compare deux versions d'un manuscrit .docx, en LECTURE SEULE.

Contrôle :
  * texte : paragraphes ajoutés, supprimés ou modifiés (avec le détail au mot près) ;
  * mise en forme visible : passages en gras / italique présents dans une seule version,
    police effective par caractère (on remonte la chaîne des styles) ;
  * champs : nombre, instructions, numéros de citation affichés ;
  * images : empreinte MD5 de chaque image, dans l'ordre ;
  * sections : marges, numérotation des lignes ; en-têtes et pieds de page.

Option --ignorer-numeros : ne pas compter comme différence les numéros de citation [n]
(utile quand seule la numérotation des références a changé).

Usage :
  python comparer_docx.py version_soumise.docx version_travail.docx
  python comparer_docx.py A.docx B.docx --ignorer-numeros > rapport.txt

Prérequis : pip install python-docx
"""
import argparse
import collections
import difflib
import hashlib
import re
import docx
from docx.oxml.ns import qn


def police_effective(p, r):
    if r.font.name:
        return r.font.name
    st = p.style
    while st is not None:
        if st.font.name:
            return st.font.name
        st = st.base_style
    return "(police par défaut / thème)"


def charger(chemin):
    d = docx.Document(chemin)
    paras = [(p.style.name, p.text) for p in d.paragraphs]
    instr, cur = [], None
    for el in d.element.body.iter():
        if el.tag == qn("w:fldChar"):
            t = el.get(qn("w:fldCharType"))
            if t == "begin":
                cur = []
            elif t == "end" and cur is not None:
                instr.append("".join(cur)); cur = None
        elif el.tag == qn("w:instrText") and cur is not None:
            cur.append(el.text or "")
    images = [hashlib.md5(d.part.related_parts[b.get(qn("r:embed"))].blob).hexdigest()
              for b in d.element.body.iter(qn("a:blip")) if b.get(qn("r:embed")) in d.part.related_parts]
    forme = {"gras": collections.Counter(), "italique": collections.Counter()}
    polices = collections.Counter()
    for p in d.paragraphs:
        for attr, cle in (("bold", "gras"), ("italic", "italique")):
            buf = ""
            for r in p.runs:
                if getattr(r, attr) or getattr(r.font, attr):
                    buf += r.text
                else:
                    if buf.strip():
                        forme[cle][buf.strip()] += 1
                    buf = ""
            if buf.strip():
                forme[cle][buf.strip()] += 1
        for r in p.runs:
            if r.text.strip():
                polices[police_effective(p, r)] += len(r.text)
    sections = [(s.page_width, s.page_height, s.left_margin, s.right_margin,
                 s._sectPr.find(qn("w:lnNumType")) is not None) for s in d.sections]
    entetes = [p.text for s in d.sections for p in list(s.header.paragraphs) + list(s.footer.paragraphs)]
    return dict(paras=paras, instr=instr, images=images, forme=forme, polices=polices, sections=sections, entetes=entetes)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("a", help="version de référence (souvent : la version réellement soumise)")
    ap.add_argument("b", help="version à comparer")
    ap.add_argument("--ignorer-numeros", action="store_true", help="ignorer les numéros de citation [n]")
    a = ap.parse_args()
    A, B = charger(a.a), charger(a.b)
    norm = (lambda s: re.sub(r"\[[\d,–\- ]+\]", "[#]", s)) if a.ignorer_numeros else (lambda s: s)

    print(f"== paragraphes : {len(A['paras'])} vs {len(B['paras'])}")
    sm = difflib.SequenceMatcher(a=[norm(t) for _, t in A["paras"]], b=[norm(t) for _, t in B["paras"]], autojunk=False)
    n = 0
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        n += 1
        if op == "replace" and i2 - i1 == j2 - j1:
            for x, y in zip(range(i1, i2), range(j1, j2)):
                ta, tb = norm(A["paras"][x][1]), norm(B["paras"][y][1])
                print(f"\n#{n} modifié A[{x}] → B[{y}]")
                for o, x1, x2, y1, y2 in difflib.SequenceMatcher(a=ta, b=tb, autojunk=False).get_opcodes():
                    if o != "equal":
                        print(f"   {o}: {ta[max(0, x1-40):x2+40]!r}\n      → {tb[max(0, y1-40):y2+40]!r}")
        else:
            for x in range(i1, i2):
                print(f"\n#{n} seulement dans A [{x}] : {A['paras'][x][1][:200]!r}")
            for y in range(j1, j2):
                print(f"\n#{n} seulement dans B [{y}] : {B['paras'][y][1][:200]!r}")
    print(f"\n== blocs de texte différents : {n}")

    print(f"== champs : {len(A['instr'])} vs {len(B['instr'])} | instructions identiques : {A['instr'] == B['instr']}")
    cit = lambda L: [m.group(1) for t in L for m in [re.search(r'"plainCitation":"([^"]*)"', t)] if m]
    print(f"== numéros de citation affichés identiques : {cit(A['instr']) == cit(B['instr'])}")
    print(f"== images : {len(A['images'])} vs {len(B['images'])} | identiques : {A['images'] == B['images']}")
    for k, (x, y) in enumerate(zip(A["images"], B["images"]), 1):
        if x != y:
            print(f"   image n°{k} différente")
    for cle in ("gras", "italique"):
        seulA, seulB = A["forme"][cle] - B["forme"][cle], B["forme"][cle] - A["forme"][cle]
        print(f"== {cle} seulement dans A : {dict(seulA) or '-'}")
        print(f"== {cle} seulement dans B : {dict(seulB) or '-'}")
    print(f"== polices (caractères) A : {dict(A['polices'])}")
    print(f"== polices (caractères) B : {dict(B['polices'])}")
    print(f"== sections / marges / numérotation des lignes identiques : {A['sections'] == B['sections']}")
    print(f"== en-têtes et pieds de page identiques : {A['entetes'] == B['entetes']}"
          + ("" if A["entetes"] == B["entetes"] else "  (souvent normal : le nombre total de pages a changé)"))
