#!/usr/bin/env python3
"""
verifier_integrite.py : compare un .docx AVANT et APRÈS modification, pour vérifier que rien
n'a été cassé en dehors des corrections voulues.

Compare :
  * le nombre de champs (Zotero et autres) et les numéros de citation affichés ;
  * les images (nombre et empreinte de chaque image, dans l'ordre) ;
  * le nombre d'entrées de bibliographie ;
  * la liste des paragraphes dont le texte a changé ;
  * les passages en italique ou en gras ajoutés ou perdus ;
  * l'empreinte (MD5) des deux fichiers.

Usage :
  python verifier_integrite.py sauvegarde.docx copie_modifiee.docx
  python verifier_integrite.py avant.docx apres.docx --style-biblio Bibliography

Prérequis : pip install python-docx
"""
import argparse
import collections
import difflib
import hashlib
import re
import docx
from docx.oxml.ns import qn


def md5(chemin):
    with open(chemin, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


def portraits(chemin, style_biblio):
    d = docx.Document(chemin)
    # reconstruire chaque instruction de champ en entier (Word la découpe parfois en plusieurs morceaux)
    instr, cur = [], None
    for el in d.element.body.iter():
        if el.tag == qn("w:fldChar"):
            t = el.get(qn("w:fldCharType"))
            if t == "begin":
                cur = []
            elif t == "end" and cur is not None:
                instr.append("".join(cur)); cur = None
        elif el.tag == qn("w:instrText") and cur is not None:
            cur.append(el.text or "")
    cits = [m.group(1) for t in instr for m in [re.search(r'"plainCitation":"([^"]*)"', t)] if m]
    debuts = sum(1 for e in d.element.body.iter(qn("w:fldChar")) if e.get(qn("w:fldCharType")) == "begin")
    images = [hashlib.md5(d.part.related_parts[b.get(qn("r:embed"))].blob).hexdigest()
              for b in d.element.body.iter(qn("a:blip")) if b.get(qn("r:embed")) in d.part.related_parts]
    paras = [p.text for p in d.paragraphs]
    biblio = sum(p.style.name == style_biblio for p in d.paragraphs)
    forme = {"gras": collections.Counter(), "italique": collections.Counter()}
    for p in d.paragraphs:
        if p.style.name == style_biblio:
            continue
        for attr, cle in (("bold", "gras"), ("italic", "italique")):
            buf = ""
            for r in p.runs:
                if getattr(r, attr) or getattr(r.font, attr):
                    buf += r.text
                elif buf.strip():
                    forme[cle][buf.strip()] += 1; buf = ""
                else:
                    buf = ""
            if buf.strip():
                forme[cle][buf.strip()] += 1
    return dict(champs=debuts, cits=cits, images=images, paras=paras, biblio=biblio, forme=forme)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("avant"); ap.add_argument("apres")
    ap.add_argument("--style-biblio", default="Bibliography")
    a = ap.parse_args()
    A, B = portraits(a.avant, a.style_biblio), portraits(a.apres, a.style_biblio)

    print(f"Empreintes : {md5(a.avant)[:10]} → {md5(a.apres)[:10]}")
    print(f"Champs : {A['champs']} → {B['champs']} | citations affichées identiques : {A['cits'] == B['cits']}")
    print(f"Images : {len(A['images'])} → {len(B['images'])} | identiques : {A['images'] == B['images']}")
    print(f"Entrées de bibliographie : {A['biblio']} → {B['biblio']}")
    sm = difflib.SequenceMatcher(a=A["paras"], b=B["paras"], autojunk=False)
    modifs = [(op, i1, i2, j1, j2) for op, i1, i2, j1, j2 in sm.get_opcodes() if op != "equal"]
    print(f"Blocs de paragraphes modifiés : {len(modifs)}")
    for op, i1, i2, j1, j2 in modifs:
        for k in range(i1, i2):
            print(f"  - avant [{k}] {A['paras'][k][:110]!r}")
        for k in range(j1, j2):
            print(f"  + après [{k}] {B['paras'][k][:110]!r}")
    for cle in ("gras", "italique"):
        perdu, gagne = A["forme"][cle] - B["forme"][cle], B["forme"][cle] - A["forme"][cle]
        if perdu or gagne:
            print(f"{cle} perdu : {dict(perdu)}")
            print(f"{cle} ajouté : {dict(gagne)}")
    print("\nÀ vérifier : tout écart non prévu (champ, image, gras/italique) doit être expliqué avant de rendre le fichier.")
